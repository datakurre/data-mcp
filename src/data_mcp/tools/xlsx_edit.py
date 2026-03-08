"""XLSX-specific sheet manipulation tools for data-mcp."""

from __future__ import annotations

from typing import Optional

import openpyxl
import pandas as pd
from fastmcp.utilities.types import ContentBlock

from data_mcp._helpers import table_response, text_response
from data_mcp.server import mcp
from data_mcp.store import get_store as _get_store


@mcp.tool
def write_sheet(
    table_name: str,
    file_path: str,
    sheet_name: str,
    if_sheet_exists: str = "replace",
    index: bool = False,
) -> list[ContentBlock]:
    """Write a table to a specific sheet in an existing XLSX file without affecting other sheets.

    Args:
        table_name: Name of the table in the store.
        file_path: Path to the target XLSX file (must already exist).
        sheet_name: Sheet name to write to.
        if_sheet_exists: Behaviour if sheet exists: 'replace' (default), 'overlay', 'new', or 'error'.
        index: Whether to include the DataFrame index (default False).
    """
    df = _get_store().get_table(table_name)
    with pd.ExcelWriter(
        file_path, engine="openpyxl", mode="a", if_sheet_exists=if_sheet_exists
    ) as writer:
        df.to_excel(writer, sheet_name=sheet_name, index=index)
    return text_response(
        f"Wrote '{table_name}' to sheet '{sheet_name}' in '{file_path}'."
    )


@mcp.tool
def manage_sheets(
    file_path: str,
    action: str = "list",
    sheet_name: Optional[str] = None,
    new_name: Optional[str] = None,
    target_sheet: Optional[str] = None,
) -> list[ContentBlock]:
    """List, delete, rename, or copy sheets within an XLSX workbook.

    Args:
        file_path: Path to the XLSX file.
        action: Operation to perform — 'list', 'delete', 'rename', or 'copy'.
        sheet_name: Sheet to operate on (required for 'delete', 'rename', and 'copy').
        new_name: New sheet name (required for 'rename').
        target_sheet: Name for the copied sheet (required for 'copy').
    """
    if action == "list":
        wb = openpyxl.load_workbook(file_path, read_only=True)
        sheets = wb.sheetnames
        wb.close()
        return text_response(f"Sheets in '{file_path}': {sheets}")

    if action == "delete":
        if not sheet_name:
            raise ValueError("'sheet_name' is required for action='delete'.")
        wb = openpyxl.load_workbook(file_path)
        if sheet_name not in wb.sheetnames:
            raise KeyError(
                f"Sheet '{sheet_name}' not found. Available: {wb.sheetnames}"
            )
        del wb[sheet_name]
        wb.save(file_path)
        wb.close()
        return text_response(f"Deleted sheet '{sheet_name}' from '{file_path}'.")

    if action == "rename":
        if not sheet_name or not new_name:
            raise ValueError(
                "'sheet_name' and 'new_name' are required for action='rename'."
            )
        wb = openpyxl.load_workbook(file_path)
        if sheet_name not in wb.sheetnames:
            raise KeyError(
                f"Sheet '{sheet_name}' not found. Available: {wb.sheetnames}"
            )
        wb[sheet_name].title = new_name
        wb.save(file_path)
        wb.close()
        return text_response(
            f"Renamed sheet '{sheet_name}' → '{new_name}' in '{file_path}'."
        )

    if action == "copy":
        if not sheet_name or not target_sheet:
            raise ValueError(
                "'sheet_name' and 'target_sheet' are required for action='copy'."
            )
        wb = openpyxl.load_workbook(file_path)
        if sheet_name not in wb.sheetnames:
            raise KeyError(
                f"Sheet '{sheet_name}' not found. Available: {wb.sheetnames}"
            )
        tgt = wb.copy_worksheet(wb[sheet_name])
        tgt.title = target_sheet
        wb.save(file_path)
        wb.close()
        return text_response(
            f"Copied sheet '{sheet_name}' → '{target_sheet}' in '{file_path}'."
        )

    raise ValueError(
        f"Unknown action '{action}'. Must be one of: 'list', 'delete', 'rename', 'copy'."
    )
