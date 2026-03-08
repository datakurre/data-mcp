"""Tests for XLSX-specific sheet manipulation tools."""

import openpyxl
import pandas as pd
import pytest

from data_mcp.store import get_store
from data_mcp.tools.xlsx_edit import manage_sheets, write_sheet


def _text(blocks) -> str:
    return blocks[0].text


@pytest.fixture
def xlsx_file(tmp_path):
    """Create an XLSX file with two sheets."""
    p = tmp_path / "workbook.xlsx"
    with pd.ExcelWriter(str(p), engine="openpyxl") as writer:
        pd.DataFrame({"a": [1, 2]}).to_excel(writer, sheet_name="Sheet1", index=False)
        pd.DataFrame({"b": [3, 4]}).to_excel(writer, sheet_name="Sheet2", index=False)
    return str(p)


def test_manage_sheets_list(xlsx_file):
    result = manage_sheets(xlsx_file, action="list")
    text = _text(result)
    assert "Sheet1" in text
    assert "Sheet2" in text


def test_manage_sheets_delete(xlsx_file):
    manage_sheets(xlsx_file, action="delete", sheet_name="Sheet2")
    wb = openpyxl.load_workbook(xlsx_file)
    assert "Sheet2" not in wb.sheetnames
    assert "Sheet1" in wb.sheetnames
    wb.close()


def test_manage_sheets_rename(xlsx_file):
    manage_sheets(xlsx_file, action="rename", sheet_name="Sheet1", new_name="Renamed")
    wb = openpyxl.load_workbook(xlsx_file)
    assert "Renamed" in wb.sheetnames
    assert "Sheet1" not in wb.sheetnames
    wb.close()


def test_manage_sheets_copy(xlsx_file):
    manage_sheets(
        xlsx_file, action="copy", sheet_name="Sheet1", target_sheet="Sheet1_copy"
    )
    wb = openpyxl.load_workbook(xlsx_file)
    assert "Sheet1_copy" in wb.sheetnames
    assert "Sheet1" in wb.sheetnames
    wb.close()


def test_manage_sheets_invalid_action(xlsx_file):
    with pytest.raises(ValueError, match="Unknown action"):
        manage_sheets(xlsx_file, action="bogus")


def test_manage_sheets_delete_missing_sheet_name(xlsx_file):
    with pytest.raises(ValueError, match="sheet_name"):
        manage_sheets(xlsx_file, action="delete")


def test_manage_sheets_rename_missing_new_name(xlsx_file):
    with pytest.raises(ValueError, match="new_name"):
        manage_sheets(xlsx_file, action="rename", sheet_name="Sheet1")


def test_manage_sheets_copy_missing_target(xlsx_file):
    with pytest.raises(ValueError, match="target_sheet"):
        manage_sheets(xlsx_file, action="copy", sheet_name="Sheet1")


def test_write_sheet(xlsx_file):
    df_new = pd.DataFrame({"x": [10, 20]})
    get_store().add_table("new_tbl", df_new)
    write_sheet("new_tbl", xlsx_file, "Sheet3")
    wb = openpyxl.load_workbook(xlsx_file)
    assert "Sheet3" in wb.sheetnames
    assert "Sheet1" in wb.sheetnames
    wb.close()
